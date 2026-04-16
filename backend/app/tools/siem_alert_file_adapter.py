"""Local ORDIV-L1A SIEM alert .xlsx adapter.

This adapter is intentionally narrow: it consumes a local .xlsx workbook,
normalizes governed ORDIV-L1A alert columns into the existing SIEM alert dict
shape, and never changes SIEMAdapterProtocol or runtime configuration.
"""

from __future__ import annotations

import asyncio
import ipaddress
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

from app.tools.siem_adapter import AdapterResult, TimeRangeSpec

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - exercised only when local env is missing openpyxl.
    load_workbook = None  # type: ignore[assignment]


GOVERNED_HEADERS: tuple[str, ...] = (
    "采集时间",
    "威胁名称",
    "威胁类型",
    "威胁等级",
    "受影响主机",
    "协议",
    "源Ip",
    "源端口",
    "目标Ip",
    "目标端口",
    "受影响主机ip",
    "状态",
    "数据来源",
    "CVE",
)

TOP_LEVEL_ALERT_KEYS = {
    "event_id",
    "event_time",
    "severity",
    "activity_name",
    "source_ip",
    "destination_ip",
    "destination_asset_id",
    "extra",
}

GAP_MISSING_REQUIRED_HEADER = "MISSING_REQUIRED_HEADER"
GAP_DUPLICATE_HEADER = "DUPLICATE_HEADER"
GAP_INVALID_WORKBOOK_SHAPE = "INVALID_WORKBOOK_SHAPE"
GAP_INVALID_TIMESTAMP = "INVALID_TIMESTAMP"
GAP_INVALID_IP_LIKE_VALUE = "INVALID_IP_LIKE_VALUE"
GAP_INVALID_PORT = "INVALID_PORT"
GAP_WORKBOOK_UNAVAILABLE = "WORKBOOK_UNAVAILABLE"
GAP_PARTIAL_ROW = "PARTIAL_ROW"

ALLOWED_GAP_CATEGORIES = {
    GAP_MISSING_REQUIRED_HEADER,
    GAP_DUPLICATE_HEADER,
    GAP_INVALID_WORKBOOK_SHAPE,
    GAP_INVALID_TIMESTAMP,
    GAP_INVALID_IP_LIKE_VALUE,
    GAP_INVALID_PORT,
    GAP_WORKBOOK_UNAVAILABLE,
    GAP_PARTIAL_ROW,
}

SEVERITY_MAP = {
    "高危": "HIGH",
    "HIGH": "HIGH",
    "中危": "MEDIUM",
    "MEDIUM": "MEDIUM",
    "低危": "LOW",
    "LOW": "LOW",
}

OUTCOME_MAP = {
    "攻击失败": "blocked",
    "疑似成功": "suspected",
}


@dataclass(frozen=True)
class _ParseOutcome:
    status: str
    alerts: list[dict[str, Any]]
    gap_reason: Optional[str]
    rows_seen: int = 0
    issues_seen: int = 0


class SIEMAlertFileAdapter:
    """ORDIV-L1A local .xlsx adapter for existing SIEM-compatible alert dicts."""

    def __init__(self, workbook_path: str | Path, *, sheet_name: str | None = None):
        self._workbook_path = Path(workbook_path)
        self._sheet_name = sheet_name
        self._last_rows_seen = 0
        self._last_alerts_returned = 0
        self._last_issues_seen = 0

    async def query_recent_summary(self, time_range: TimeRangeSpec) -> AdapterResult[dict]:
        started = time.monotonic()
        parsed = await self._parse_alerts()
        alerts = self._filter_alerts_by_time(parsed.alerts, time_range)
        severity_counts: dict[str, int] = {}
        for alert in alerts:
            severity = str(alert.get("severity") or "UNKNOWN")
            severity_counts[severity] = severity_counts.get(severity, 0) + 1

        summary = {
            "total_alerts": len(alerts),
            "severity_counts": severity_counts,
        }
        return self._adapter_result(parsed, summary, started)

    async def query_asset_alerts(self, asset_id: str, time_range: TimeRangeSpec) -> AdapterResult[list]:
        started = time.monotonic()
        parsed = await self._parse_alerts()
        alerts = self._filter_alerts_by_time(parsed.alerts, time_range)
        filtered = [
            alert
            for alert in alerts
            if asset_id
            and (
                alert.get("destination_asset_id") == asset_id
                or alert.get("source_ip") == asset_id
                or alert.get("destination_ip") == asset_id
            )
        ]
        return self._adapter_result(parsed, filtered, started)

    async def query_intent_alerts(
        self,
        intent: str,
        user_input: str,
        time_range: TimeRangeSpec,
    ) -> AdapterResult[list]:
        started = time.monotonic()
        parsed = await self._parse_alerts()
        alerts = self._filter_alerts_by_time(parsed.alerts, time_range)
        return self._adapter_result(parsed, alerts, started)

    async def get_scenario_metadata(self, scenario_id: str) -> AdapterResult[Optional[dict]]:
        return AdapterResult.ok(None)

    async def get_asset_context(self, asset_id: str) -> AdapterResult[Optional[dict]]:
        return AdapterResult.ok(None)

    def get_runtime_stats(self) -> dict[str, int]:
        return {
            "rows_seen": self._last_rows_seen,
            "alerts_returned": self._last_alerts_returned,
            "issues_seen": self._last_issues_seen,
        }

    async def _parse_alerts(self) -> _ParseOutcome:
        parsed = await asyncio.to_thread(self._parse_workbook)
        self._last_rows_seen = parsed.rows_seen
        self._last_alerts_returned = len(parsed.alerts)
        self._last_issues_seen = parsed.issues_seen
        return parsed

    def _parse_workbook(self) -> _ParseOutcome:
        if load_workbook is None:
            return _ParseOutcome("unavailable", [], GAP_WORKBOOK_UNAVAILABLE, issues_seen=1)

        if self._workbook_path.suffix.lower() != ".xlsx":
            return _ParseOutcome("unavailable", [], GAP_WORKBOOK_UNAVAILABLE, issues_seen=1)

        workbook = None
        try:
            workbook = load_workbook(self._workbook_path, read_only=True, data_only=True)
        except Exception:
            return _ParseOutcome("unavailable", [], GAP_WORKBOOK_UNAVAILABLE, issues_seen=1)

        try:
            worksheet = self._select_worksheet(workbook)
            if worksheet is None:
                return _ParseOutcome("partial", [], GAP_INVALID_WORKBOOK_SHAPE, issues_seen=1)

            header_row_number, header_map, header_gap = self._find_header_row(worksheet)
            if header_gap:
                return _ParseOutcome("partial", [], header_gap, issues_seen=1)
            if header_row_number is None or header_map is None:
                return _ParseOutcome("partial", [], GAP_INVALID_WORKBOOK_SHAPE, issues_seen=1)

            alerts: list[dict[str, Any]] = []
            issues: list[str] = []
            rows_seen = 0
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=header_row_number + 1, values_only=True),
                start=header_row_number + 1,
            ):
                if self._is_blank_row(row):
                    continue
                rows_seen += 1
                alert, row_issues = self._parse_alert_row(row, row_number, header_map)
                issues.extend(row_issues)
                if alert is not None:
                    alerts.append(alert)

            first_issue = issues[0] if issues else None
            return _ParseOutcome(
                "partial" if first_issue else "ok",
                alerts,
                first_issue,
                rows_seen=rows_seen,
                issues_seen=len(issues),
            )
        finally:
            if workbook is not None:
                workbook.close()

    def _select_worksheet(self, workbook: Any) -> Any | None:
        if not getattr(workbook, "worksheets", None):
            return None
        if self._sheet_name is None:
            return workbook.active
        if self._sheet_name not in workbook.sheetnames:
            return None
        return workbook[self._sheet_name]

    def _find_header_row(self, worksheet: Any) -> tuple[int | None, dict[str, int] | None, Optional[str]]:
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            values = [_safe_text(value) for value in row]
            positions: dict[str, int] = {}
            duplicate_header = False
            for index, value in enumerate(values):
                if value not in GOVERNED_HEADERS:
                    continue
                if value in positions:
                    duplicate_header = True
                positions.setdefault(value, index)

            if duplicate_header:
                return None, None, GAP_DUPLICATE_HEADER
            if all(header in positions for header in GOVERNED_HEADERS):
                return row_number, positions, None

        return None, None, GAP_MISSING_REQUIRED_HEADER

    def _parse_alert_row(
        self,
        row: tuple[Any, ...],
        row_number: int,
        header_map: dict[str, int],
    ) -> tuple[dict[str, Any] | None, list[str]]:
        values = {
            header: _safe_text(row[index] if index < len(row) else None)
            for header, index in header_map.items()
        }

        if not values.get("采集时间") or not values.get("威胁名称") or not values.get("威胁等级"):
            return None, [_gap(GAP_PARTIAL_ROW, row_number)]

        event_time = _parse_event_time(row[header_map["采集时间"]])
        if event_time is None:
            return None, [_gap(GAP_INVALID_TIMESTAMP, row_number)]

        severity = _normalize_severity(values["威胁等级"])
        if severity is None:
            return None, [_gap(GAP_PARTIAL_ROW, row_number)]

        source_ip = _normalize_ip(values.get("源Ip", ""))
        destination_ip = _normalize_ip(values.get("目标Ip", ""))
        affected_host_ip = _normalize_ip(values.get("受影响主机ip", ""))
        if source_ip is None or destination_ip is None or affected_host_ip is None:
            return None, [_gap(GAP_INVALID_IP_LIKE_VALUE, row_number)]

        issues: list[str] = []
        source_port = _normalize_port(values.get("源端口", ""))
        destination_port = _normalize_port(values.get("目标端口", ""))
        if source_port is None and values.get("源端口"):
            issues.append(_gap(GAP_INVALID_PORT, row_number))
        if destination_port is None and values.get("目标端口"):
            issues.append(_gap(GAP_INVALID_PORT, row_number))

        outcome = OUTCOME_MAP.get(values.get("状态", ""), "unknown")
        extra = {
            "alert_name": values["威胁名称"],
            "threat_category": values.get("威胁类型", ""),
            "affected_host_label": values.get("受影响主机", ""),
            "protocol": values.get("协议", ""),
            "affected_host_ip": affected_host_ip,
            "outcome": outcome,
            "detection_engine": values.get("数据来源", ""),
            "cve_id": values.get("CVE", ""),
        }
        if source_port is not None:
            extra["source_port"] = source_port
        if destination_port is not None:
            extra["destination_port"] = destination_port

        alert = {
            "event_id": f"ordiv_l1a_row_{row_number}",
            "event_time": _format_iso_z(event_time),
            "severity": severity,
            "activity_name": values["威胁名称"],
            "source_ip": source_ip,
            "destination_ip": destination_ip,
            "destination_asset_id": values.get("受影响主机", ""),
            "extra": extra,
        }
        return {key: alert[key] for key in TOP_LEVEL_ALERT_KEYS}, issues

    @staticmethod
    def _is_blank_row(row: tuple[Any, ...]) -> bool:
        return all(_safe_text(value) == "" for value in row)

    @staticmethod
    def _filter_alerts_by_time(alerts: list[dict[str, Any]], time_range: TimeRangeSpec) -> list[dict[str, Any]]:
        filtered: list[dict[str, Any]] = []
        for alert in alerts:
            event_time = _parse_iso_z(alert.get("event_time", ""))
            if time_range.contains(event_time):
                filtered.append(alert)
        return filtered

    @staticmethod
    def _adapter_result(parsed: _ParseOutcome, data: Any, started: float) -> AdapterResult[Any]:
        latency_ms = (time.monotonic() - started) * 1000
        metadata = {
            "rows_seen": parsed.rows_seen,
            "alerts_returned": len(data) if isinstance(data, list) else len(parsed.alerts),
            "issues_seen": parsed.issues_seen,
        }
        if parsed.status == "unavailable":
            return AdapterResult.unavailable(
                data,
                gap_reason=parsed.gap_reason or GAP_WORKBOOK_UNAVAILABLE,
                latency_ms=latency_ms,
                metadata=metadata,
            )
        if parsed.status == "partial":
            return AdapterResult.partial(
                data,
                gap_reason=parsed.gap_reason or GAP_PARTIAL_ROW,
                latency_ms=latency_ms,
                metadata=metadata,
            )
        return AdapterResult.ok(data, latency_ms=latency_ms, metadata=metadata)


def _safe_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).strip().split())


def _gap(category: str, row_number: int | None = None) -> str:
    if category not in ALLOWED_GAP_CATEGORIES:
        category = GAP_PARTIAL_ROW
    if row_number is None:
        return category
    return f"{category}:row_{row_number}"


def _parse_event_time(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        try:
            parsed = datetime.strptime(value.strip(), "%Y-%m-%d %H:%M:%S")
        except ValueError:
            return None
    else:
        return None

    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def _parse_iso_z(value: str) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(timezone.utc)
    except ValueError:
        return None


def _format_iso_z(value: datetime) -> str:
    return value.astimezone(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _normalize_severity(value: str) -> str | None:
    token = _safe_text(value)
    return SEVERITY_MAP.get(token) or SEVERITY_MAP.get(token.upper())


def _normalize_ip(value: str) -> str | None:
    token = _safe_text(value)
    if not token:
        return ""
    try:
        return str(ipaddress.ip_address(token))
    except ValueError:
        return None


def _normalize_port(value: str) -> str | None:
    token = _safe_text(value)
    if not token:
        return ""
    if not token.isdigit():
        return None
    port = int(token)
    if 1 <= port <= 65535:
        return str(port)
    return None
